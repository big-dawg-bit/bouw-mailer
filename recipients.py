"""Leest het Excel-bestand en levert een ontdubbelde lijst ontvangers."""

import re
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from config import (
    CATEGORY_VARIANTS,
    COL_CATEGORY,
    COL_CITY,
    COL_COMPANY,
    COL_EMAIL,
    SHEET_NAME,
    VARIANT_PRIORITY,
)

# Bewust streng: we mailen liever een adres te weinig dan een kapot adres.
EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$")


@dataclass
class Recipient:
    email: str
    company: str
    category: str
    city: str
    variant: str
    # Andere bedrijven in het bestand met hetzelfde adres.
    also: list


@dataclass
class LoadReport:
    recipients: list
    rows_read: int
    skipped_no_email: int
    skipped_bad_email: list
    skipped_unknown_category: list
    merged: list


def _headers(sheet) -> dict:
    """Kolomindex per kop. Op naam, niet op letter: het bestand wordt
    periodiek opnieuw gegenereerd en kolommen kunnen verschuiven."""
    first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    found = {
        str(name).strip(): idx
        for idx, name in enumerate(first)
        if name is not None and str(name).strip()
    }
    missing = [c for c in (COL_COMPANY, COL_CATEGORY, COL_EMAIL) if c not in found]
    if missing:
        sys.exit(
            "FOUT: deze kolommen ontbreken in het werkblad: "
            + ", ".join(missing)
            + "\n      Gevonden koppen: "
            + ", ".join(sorted(found))
        )
    return found


def load_recipients(excel_path: Path) -> LoadReport:
    if not excel_path.exists():
        sys.exit(
            f"FOUT: Excel-bestand niet gevonden: {excel_path}\n"
            f"      Zet het bestand in data/ of pas EXCEL_PATH in .env aan."
        )

    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        sys.exit(
            f"FOUT: werkblad '{SHEET_NAME}' niet gevonden. "
            f"Aanwezig: {', '.join(workbook.sheetnames)}"
        )
    sheet = workbook[SHEET_NAME]
    cols = _headers(sheet)

    def cell(row, key, default=""):
        idx = cols.get(key)
        if idx is None or idx >= len(row) or row[idx] is None:
            return default
        return str(row[idx]).strip()

    rows_read = 0
    skipped_no_email = 0
    skipped_bad_email = []
    skipped_unknown_category = []
    # adres (lowercase) -> Recipient; botsingen lossen we hieronder op
    by_email = {}
    merged = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        company = cell(row, COL_COMPANY)
        if not company:
            continue
        rows_read += 1

        email = cell(row, COL_EMAIL)
        if not email:
            skipped_no_email += 1
            continue

        if not EMAIL_RE.match(email):
            skipped_bad_email.append((company, email))
            continue

        category = cell(row, COL_CATEGORY)
        variant = CATEGORY_VARIANTS.get(category.lower())
        if not variant:
            skipped_unknown_category.append((company, category or "(leeg)"))
            continue

        key = email.lower()
        existing = by_email.get(key)
        if existing is None:
            by_email[key] = Recipient(
                email=email,
                company=company,
                category=category,
                city=cell(row, COL_CITY),
                variant=variant,
                also=[],
            )
            continue

        # Zelfde adres, ander bedrijf: één mail sturen, niet twee.
        existing.also.append(company)
        if VARIANT_PRIORITY.index(variant) < VARIANT_PRIORITY.index(existing.variant):
            merged.append(
                f"{email}: {company} ({variant}) wint van "
                f"{existing.company} ({existing.variant})"
            )
            existing.variant = variant
            existing.category = category
        else:
            merged.append(
                f"{email}: {company} ({variant}) valt samen met "
                f"{existing.company} ({existing.variant})"
            )

    workbook.close()

    recipients = sorted(
        by_email.values(),
        key=lambda r: (VARIANT_PRIORITY.index(r.variant), r.company.lower()),
    )

    return LoadReport(
        recipients=recipients,
        rows_read=rows_read,
        skipped_no_email=skipped_no_email,
        skipped_bad_email=skipped_bad_email,
        skipped_unknown_category=skipped_unknown_category,
        merged=merged,
    )
